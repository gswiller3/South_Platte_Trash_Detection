"""
Build the field results workbook from the field_data folder.

Walks field_data/, finds every camera_XX folder that process_camera.py wrote, reads
its detections.csv, summary.txt and review/validation_result.txt, and writes one
Excel workbook with the observed counts per camera, per 14 day period and per day.
The Comparison sheet has the same columns as the one in
South_Platte_Predicted_Trash_Dataset.xlsx, with the ACTUAL cells filled in by this
script instead of by hand, and the predicted low, mid and high pulled from that
workbook when it is in the same folder.

Usage:
    python build_field_workbook.py
    python build_field_workbook.py --field-dir field_data --out South_Platte_Field_Results.xlsx
    python build_field_workbook.py --predicted path/to/South_Platte_Predicted_Trash_Dataset.xlsx

Run it again after every card. It rebuilds the whole workbook from the folders, so
nothing is typed in by hand and nothing gets out of sync.

The camera clocks were set a year behind, so timestamps on the cards say 2025. The
script rewrites the year to STUDY_YEAR before sorting frames into periods, and the
Cameras sheet says whether that happened for each camera.

Author: Griffin Willer
"""

import argparse
import csv
import math
import os
import re
from datetime import date, datetime, timedelta
from decimal import Decimal, ROUND_HALF_UP

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# CONFIG, keep in step with build_prediction_dataset.py
# ---------------------------------------------------------------------------

STUDY_YEAR = 2026
PERIODS = [
    ("P1", date(2026, 9, 28), date(2026, 10, 11)),
    ("P2", date(2026, 10, 12), date(2026, 10, 25)),
    ("P3", date(2026, 10, 26), date(2026, 11, 8)),
]
FRAME_INTERVAL_S = 10          # time lapse interval the cameras were set to
DENVER_LAT = 39.74
LOW_MULT, HIGH_MULT = 0.25, 3.0   # only used if the predicted workbook cannot be read

PREDICTED_XLSX = "South_Platte_Predicted_Trash_Dataset.xlsx"
PREDICTED_PERIOD_CSV = "heatmap_periods_long.csv"
PREDICTED_DAILY_CSV = "heatmap_daily_long.csv"

# Camera number to bridge. Same names and coordinates as the predicted dataset so the
# two line up. Coordinates are National Bridge Inventory positions, replace with the
# GPS taken at install when the pin to bridge mapping is confirmed.
SITES = [
    dict(id=1, bridge="York St (Commerce City, nearest to E 56th Ave)", lat=39.80106, lon=-104.95935),
    dict(id=2, bridge="E 51st Ave / Bettie Cram Dr (National Western Center) or 47th Ave", lat=39.78851, lon=-104.97271),
    dict(id=3, bridge="31st St (Arkins Ct, Denargo rail yards) or 38th St", lat=39.76766, lon=-104.98496),
    dict(id=4, bridge="19th St / 20th St (Commons Park north end)", lat=39.76050, lon=-105.00365),
    dict(id=5, bridge="15th St (Confluence Park) or 16th St pedestrian bridge", lat=39.75550, lon=-105.00788),
    dict(id=6, bridge="Speer Blvd", lat=39.75443, lon=-105.01026),
    dict(id=7, bridge="W Colfax Ave (Auraria)", lat=39.73998, lon=-105.01789),
    dict(id=8, bridge="W 8th Ave (Sun Valley, Frog Hollow Park) or US 6 (6th Ave)", lat=39.72940, lon=-105.01717),
    dict(id=9, bridge="W Alameda Ave", lat=39.71107, lon=-105.00036),
    dict(id=10, bridge="W Mississippi Ave", lat=39.69667, lon=-104.99249),
    dict(id=11, bridge="W Florida Ave (Ruby Hill Park) or W Evans Ave", lat=39.68912, lon=-104.99982),
]
SITE_BY_ID = {s["id"]: s for s in SITES}

# ---------------------------------------------------------------------------
# STYLES, same as the predicted workbook
# ---------------------------------------------------------------------------

HDR_FILL = PatternFill("solid", fgColor="1F3A5F")
HDR_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name="Arial", size=10)
BOLD_FONT = Font(name="Arial", bold=True, size=10)
NOTE_FONT = Font(name="Arial", size=9, italic=True, color="555555")
GREEN = PatternFill("solid", fgColor="E2EFDA")
THIN = Side(style="thin", color="BBBBBB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def style_body(ws, row, ncols, font=BODY_FONT):
    for c in range(1, ncols + 1):
        ws.cell(row=row, column=c).font = font
        ws.cell(row=row, column=c).border = BORDER


def autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_table(ws, headers, rows, widths=None, freeze="A2"):
    for c, h in enumerate(headers, start=1):
        ws.cell(row=1, column=c, value=h)
    style_header(ws, 1, len(headers))
    for r, row in enumerate(rows, start=2):
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v)
        style_body(ws, r, len(headers))
    if widths:
        autosize(ws, widths)
    ws.freeze_panes = freeze
    return len(rows) + 1


# ---------------------------------------------------------------------------
# HELPERS
# ---------------------------------------------------------------------------

def daylight_hours(d):
    n = d.timetuple().tm_yday
    decl = 23.44 * math.sin(math.radians(360.0 / 365.0 * (n - 81)))
    x = -math.tan(math.radians(DENVER_LAT)) * math.tan(math.radians(decl))
    x = max(-1.0, min(1.0, x))
    return 2.0 * math.degrees(math.acos(x)) / 15.0


def period_daylight():
    out = {}
    for name, s, e in PERIODS:
        dd, tot = s, 0.0
        while dd <= e:
            tot += daylight_hours(dd)
            dd += timedelta(days=1)
        out[name] = tot
    return out


def period_of(d):
    for name, s, e in PERIODS:
        if s <= d <= e:
            return name
    return "outside"


def fix_year(ts, study_year):
    if ts.year == study_year:
        return ts, False
    try:
        return ts.replace(year=study_year), True
    except ValueError:        # Feb 29 on a non leap year, never happens in this window
        return ts.replace(year=study_year, day=28), True


def r2(x):
    # round half up like Excel does, so the CSV and the sheet show the same number
    return float(Decimal(str(x)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def median(vals):
    if not vals:
        return None
    v = sorted(vals)
    n = len(v)
    return v[n // 2] if n % 2 else (v[n // 2 - 1] + v[n // 2]) / 2.0


# ---------------------------------------------------------------------------
# READ ONE CAMERA FOLDER
# ---------------------------------------------------------------------------

def read_summary(path):
    out = {"setup_skipped": None, "videos": None, "conf": None, "imgsz": None, "tiles": None, "model": None}
    if not os.path.exists(path):
        return out
    with open(path) as f:
        for line in f:
            line = line.strip()
            if line.startswith("Setup frames skipped:"):
                out["setup_skipped"] = int(line.split(":")[1])
            elif line.startswith("Videos skipped:"):
                out["videos"] = int(line.split(":")[1])
            elif line.startswith("Confidence threshold:"):
                m = re.search(r"Confidence threshold:\s*([\d.]+)", line)
                out["conf"] = float(m.group(1)) if m else None
                m = re.search(r"Inference size:\s*(\d+)", line)
                out["imgsz"] = int(m.group(1)) if m else None
                m = re.search(r"Tiles:\s*(\d+)", line)
                out["tiles"] = int(m.group(1)) if m else 1
                m = re.search(r"Model:\s*(\S+)", line)
                out["model"] = m.group(1) if m else None
    return out


def read_validation(path):
    if not os.path.exists(path):
        return None
    out = {}
    with open(path) as f:
        for line in f:
            parts = line.strip().split(",")
            if len(parts) == 2:
                out[parts[0]] = int(parts[1])
            elif len(parts) == 4:
                out[parts[0]] = tuple(float(x) for x in parts[1:])
    return out if "precision" in out else None


def read_camera(cam_dir, study_year, interval):
    name = os.path.basename(cam_dir)
    m = re.match(r"camera_(\d+)$", name)
    if not m:
        return None
    cam_id = int(m.group(1))
    det_path = os.path.join(cam_dir, "detections.csv")
    if not os.path.exists(det_path):
        print("  %s has no detections.csv, skipped" % name)
        return None

    frames = {}          # filename -> frame record
    boxes = []           # box rows for the Detections sheet
    year_fixed = 0
    with open(det_path) as f:
        for row in csv.DictReader(f):
            fname = row["filename"]
            ts = datetime.strptime(row["timestamp"], "%Y-%m-%d %H:%M:%S")
            ts, fixed = fix_year(ts, study_year)
            if fixed:
                year_fixed += 1
            idx = int(row["box_index"])
            fr = frames.get(fname)
            if fr is None:
                fr = dict(filename=fname, ts=ts, date=ts.date(), period=period_of(ts.date()),
                          night=(idx == -2), boxes=0, new=0, stationary=0, drift=0, max_conf=0.0)
                frames[fname] = fr
            if idx >= 0:
                conf = float(row["confidence"])
                rep = row.get("repeat_type", "") or ("new" if row.get("is_new") == "1" else "repeat")
                fr["boxes"] += 1
                fr[rep] = fr.get(rep, 0) + 1
                fr["max_conf"] = max(fr["max_conf"], conf)
                boxes.append(dict(camera_id=cam_id, filename=fname, ts=ts, period=fr["period"], box_index=idx,
                                  x1=row["x1"], y1=row["y1"], x2=row["x2"], y2=row["y2"],
                                  confidence=conf, repeat_type=rep))

    recs = sorted(frames.values(), key=lambda r: (r["ts"], r["filename"]))
    processed = [r for r in recs if not r["night"]]
    gaps = []
    for a, b in zip(processed, processed[1:]):
        g = (b["ts"] - a["ts"]).total_seconds()
        if 0 < g < 3600:
            gaps.append(g)

    summ = read_summary(os.path.join(cam_dir, "summary.txt"))
    val = read_validation(os.path.join(cam_dir, "review", "validation_result.txt"))

    return dict(id=cam_id, folder=cam_dir, frames=recs, processed=processed, boxes=boxes,
                night=sum(1 for r in recs if r["night"]),
                year_fixed=year_fixed, median_gap=median(gaps), summary=summ, validation=val, interval=interval)


def bucket(cam, key_fn):
    # Sum frame stats by a key such as period or date.
    out = {}
    for r in cam["processed"]:
        k = key_fn(r)
        b = out.setdefault(k, dict(frames=0, with_trash=0, boxes=0, new=0, stationary=0, drift=0, days=set()))
        b["frames"] += 1
        b["with_trash"] += 1 if r["boxes"] else 0
        b["boxes"] += r["boxes"]
        b["new"] += r.get("new", 0)
        b["stationary"] += r.get("stationary", 0)
        b["drift"] += r.get("drift", 0)
        b["days"].add(r["date"])
    return out


# ---------------------------------------------------------------------------
# PREDICTED NUMBERS
# ---------------------------------------------------------------------------

def read_predicted(xlsx_path):
    # Returns ({(pin, period): (mid, low, high)}, source description)
    pred = {}
    if xlsx_path and os.path.exists(xlsx_path):
        try:
            wb = load_workbook(xlsx_path, data_only=True)
            ws = wb["Predicted_Counts"]
            for row in ws.iter_rows(min_row=2, max_col=15, values_only=True):
                pin = row[0]
                if not isinstance(pin, int):
                    continue
                for k, (name, _, _) in enumerate(PERIODS):
                    mid, lo, hi = row[5 + k], row[9 + k], row[12 + k]
                    if isinstance(mid, (int, float)) and isinstance(lo, (int, float)) and isinstance(hi, (int, float)):
                        pred[(pin, name)] = (mid, lo, hi)
            if pred:
                return pred, os.path.basename(xlsx_path)
        except Exception as e:
            print("Could not read predicted workbook %s: %s" % (xlsx_path, e))
    # Fallback: the period CSV that build_prediction_dataset.py writes next to the workbook
    for folder in (os.path.dirname(xlsx_path or "") or ".", "."):
        p = os.path.join(folder, PREDICTED_PERIOD_CSV)
        if os.path.exists(p):
            with open(p) as f:
                for row in csv.DictReader(f):
                    mid = float(row["predicted_new_items"])
                    pred[(int(row["site_id"]), row["period"])] = (mid, round(mid * LOW_MULT), round(mid * HIGH_MULT))
            if pred:
                return pred, os.path.basename(p) + " (low and high = mid x %.2f and x %.1f)" % (LOW_MULT, HIGH_MULT)
    return {}, None


def read_predicted_daily(xlsx_path):
    # {(pin, date): (predicted_new_items, flow_cfs)} from heatmap_daily_long.csv if present
    out = {}
    for folder in (os.path.dirname(xlsx_path or "") or ".", "."):
        p = os.path.join(folder, PREDICTED_DAILY_CSV)
        if os.path.exists(p):
            with open(p) as f:
                for row in csv.DictReader(f):
                    d = datetime.strptime(row["date"], "%Y-%m-%d").date()
                    flow = row.get("flow_cfs", "")
                    out[(int(row["site_id"]), d)] = (float(row["predicted_new_items"]), float(flow) if flow else None)
            return out
    return out


# ---------------------------------------------------------------------------
# WORKBOOK
# ---------------------------------------------------------------------------

def build_workbook(cams, pred, pred_src, pred_daily, interval, out_dir):
    wb = Workbook()
    dl_sum = period_daylight()
    today = date.today().isoformat()
    n_boxes_total = sum(len(c["boxes"]) for c in cams)

    # ---------------- README ----------------
    ws = wb.active
    ws.title = "README"
    lines = [
        ("South Platte River Trash Source Study", True),
        ("Field results workbook, built %s by build_field_workbook.py from %d camera folders" % (today, len(cams)), False),
        ("These are MEASURED counts from the bridge cameras, run through best_river.pt by process_camera.py.", False),
        ("", False),
        ("How it was built", True),
        ("Every camera_XX folder under field_data was read. detections.csv gives one row per box with its timestamp and", False),
        ("repeat type, plus one row per frame with no boxes and one row per night frame that was skipped.", False),
        ("Frames were sorted into the three 14 day periods by date: P1 Sep 28 to Oct 11, P2 Oct 12 to Oct 25, P3 Oct 26 to Nov 8.", False),
        ("The camera clocks were set a year behind, so any timestamp not in %d had its year rewritten. The Cameras sheet shows where that happened." % STUDY_YEAR, False),
        ("New items = boxes tagged new by process_camera.py, which is the count after stationary and drift repeats are removed.", False),
        ("Daylight hours covered = daylight frames processed x %d s. This is the real coverage, not the 14 day period, so compare rates per hour." % interval, False),
        ("Predicted numbers on Comparison and Daily come from: %s" % (pred_src or "not found, put South_Platte_Predicted_Trash_Dataset.xlsx next to this script and rebuild"), False),
        ("", False),
        ("Sheets", True),
        ("Cameras: one row per camera, the summary.txt numbers plus coverage, rate per daylight hour and validation if scored.", False),
        ("By_Period: camera by period, new items, frames, hours and rate. Same shape as Predicted_Counts in the predicted workbook.", False),
        ("Comparison: predicted low, mid, high next to the observed count for every camera and period. Filled by the script.", False),
        ("Heatmap_Long: one row per camera per period with lat, lon and weight. Load this into the heat map tool. Also saved as heatmap_observed_long.csv.", False),
        ("Daily: one row per camera per day with frames, hours, new items, the predicted number and the gauge flow for that day.", False),
        ("Validation: precision, miss rate and recall from score_review.py for every camera that has a scored review sheet.", False),
        ("Detections: every box from every camera, %d rows. Filter by repeat_type = new for the counted items." % n_boxes_total, False),
        ("", False),
        ("Reading the numbers", True),
        ("Report New items, not Total boxes. Total boxes counts the same piece of trash every frame it sits in.", False),
        ("A camera with no frames in a period shows blanks on Comparison, not zeros. Zero means the camera ran and saw nothing.", False),
        ("Rates per daylight hour are the fair comparison between cameras, since no two cards cover the same hours.", False),
        ("Validation precision is image level. New items x precision on the Cameras sheet is a rough corrected count, not a measured one.", False),
        ("Rebuild this workbook after every card: python3 build_field_workbook.py", False),
    ]
    for i, (t, bold) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=t)
        c.font = Font(name="Arial", size=11 if bold else 10, bold=bold)
    ws.column_dimensions["A"].width = 130

    # ---------------- Cameras ----------------
    wcam = wb.create_sheet("Cameras")
    hdr = ["Pin", "Bridge", "First frame", "Last frame", "Frames in file", "Night frames skipped", "Setup frames skipped",
           "Daylight frames processed", "Daylight hours covered", "Median gap between frames (s)", "Images with trash",
           "Total trash boxes", "Stationary repeats", "Drift repeats", "New items", "New items per daylight h",
           "Validated?", "Precision", "Recall", "New items x precision", "Clock year fixed?", "Conf", "Inference size", "Tiles", "Model", "Folder"]
    for c, h in enumerate(hdr, start=1):
        wcam.cell(row=1, column=c, value=h)
    style_header(wcam, 1, len(hdr))
    for i, cam in enumerate(cams, start=2):
        site = SITE_BY_ID.get(cam["id"], {})
        s = cam["summary"]
        tot = bucket(cam, lambda r: "all").get("all", dict(frames=0, with_trash=0, boxes=0, new=0, stationary=0, drift=0))
        val = cam["validation"]
        first = min(r["ts"] for r in cam["frames"]) if cam["frames"] else None
        last = max(r["ts"] for r in cam["frames"]) if cam["frames"] else None
        row = [cam["id"], site.get("bridge", "pin %d not in site table" % cam["id"]),
               first.strftime("%Y-%m-%d %H:%M") if first else "", last.strftime("%Y-%m-%d %H:%M") if last else "",
               len(cam["frames"]), cam["night"], s["setup_skipped"] if s["setup_skipped"] is not None else "",
               tot["frames"], "=ROUND(H%d*%d/3600,2)" % (i, interval), round(cam["median_gap"], 1) if cam["median_gap"] else "",
               tot["with_trash"], tot["boxes"], tot["stationary"], tot["drift"], tot["new"],
               "=IF(H%d=0,\"\",ROUND(O%d/(H%d*%d/3600),2))" % (i, i, i, interval),
               "yes" if val else "no",
               round(val["precision"][0], 3) if val else "", round(val["recall"][0], 3) if val else "",
               "=IF(R%d=\"\",\"\",ROUND(O%d*R%d,0))" % (i, i, i) if val else "",
               "yes, %d timestamps" % cam["year_fixed"] if cam["year_fixed"] else "no",
               s["conf"] if s["conf"] is not None else "", s["imgsz"] if s["imgsz"] is not None else "",
               s["tiles"] if s["tiles"] is not None else "", s["model"] or "", cam["folder"]]
        for c, v in enumerate(row, start=1):
            wcam.cell(row=i, column=c, value=v)
        style_body(wcam, i, len(hdr))
    tr = len(cams) + 2
    wcam.cell(row=tr, column=2, value="Corridor total").font = BOLD_FONT
    for c in (5, 6, 8, 11, 12, 13, 14, 15):
        L = get_column_letter(c)
        wcam.cell(row=tr, column=c, value="=SUM(%s2:%s%d)" % (L, L, tr - 1)).font = BOLD_FONT
    wcam.cell(row=tr, column=9, value="=ROUND(H%d*%d/3600,2)" % (tr, interval)).font = BOLD_FONT
    wcam.cell(row=tr, column=16, value="=IF(H%d=0,\"\",ROUND(O%d/(H%d*%d/3600),2))" % (tr, tr, tr, interval)).font = BOLD_FONT
    wcam.cell(row=tr + 2, column=1, value="New items is the number to report. Daylight hours covered = daylight frames x %d s. Median gap should be about %d s; if it is much larger the camera was on motion trigger or dropped frames." % (interval, interval)).font = NOTE_FONT
    wcam.cell(row=tr + 3, column=1, value="Precision and recall come from score_review.py and are image level. New items x precision is a rough corrected count for the report, say so when you use it.").font = NOTE_FONT
    autosize(wcam, [6, 40, 16, 16, 10, 10, 10, 12, 12, 12, 10, 10, 10, 10, 10, 12, 10, 10, 10, 12, 16, 7, 10, 7, 18, 30])
    wcam.freeze_panes = "C2"

    # ---------------- By_Period ----------------
    wp = wb.create_sheet("By_Period")
    pnames = [p[0] for p in PERIODS]
    hdr = ["Pin", "Bridge", "Lat", "Lon"] + ["%s new items" % p for p in pnames] + ["Total new items", "Outside window new items"] + \
          ["%s daylight frames" % p for p in pnames] + ["%s daylight hours" % p for p in pnames] + \
          ["%s items per daylight h" % p for p in pnames] + ["%s raw boxes" % p for p in pnames] + ["%s days with frames" % p for p in pnames]
    for c, h in enumerate(hdr, start=1):
        wp.cell(row=1, column=c, value=h)
    style_header(wp, 1, len(hdr))
    per = {}   # (pin, period) -> bucket
    for cam in cams:
        for k, b in bucket(cam, lambda r: r["period"]).items():
            per[(cam["id"], k)] = b
    for i, cam in enumerate(cams, start=2):
        site = SITE_BY_ID.get(cam["id"], {})
        wp.cell(row=i, column=1, value=cam["id"])
        wp.cell(row=i, column=2, value=site.get("bridge", ""))
        wp.cell(row=i, column=3, value=site.get("lat", ""))
        wp.cell(row=i, column=4, value=site.get("lon", ""))
        for k, name in enumerate(pnames):
            b = per.get((cam["id"], name))
            wp.cell(row=i, column=5 + k, value=b["new"] if b else "")
            wp.cell(row=i, column=10 + k, value=b["frames"] if b else "")
            fcol = get_column_letter(10 + k)
            wp.cell(row=i, column=13 + k, value="=IF(%s%d=\"\",\"\",ROUND(%s%d*%d/3600,2))" % (fcol, i, fcol, i, interval))
            ncol = get_column_letter(5 + k)
            wp.cell(row=i, column=16 + k, value="=IF(OR(%s%d=\"\",%s%d=0),\"\",ROUND(%s%d/(%s%d*%d/3600),2))" % (ncol, i, fcol, i, ncol, i, fcol, i, interval))
            wp.cell(row=i, column=19 + k, value=b["boxes"] if b else "")
            wp.cell(row=i, column=22 + k, value=len(b["days"]) if b else "")
        wp.cell(row=i, column=8, value="=SUM(E%d:G%d)" % (i, i))
        ob = per.get((cam["id"], "outside"))
        wp.cell(row=i, column=9, value=ob["new"] if ob else "")
        style_body(wp, i, len(hdr))
    tr = len(cams) + 2
    wp.cell(row=tr, column=2, value="Corridor total").font = BOLD_FONT
    for c in list(range(5, 13)) + list(range(19, 22)):
        L = get_column_letter(c)
        wp.cell(row=tr, column=c, value="=SUM(%s2:%s%d)" % (L, L, tr - 1)).font = BOLD_FONT
    wp.cell(row=tr + 2, column=1, value="Blank = the camera has no daylight frames in that period. Outside window = frames dated before Sep 28 or after Nov 8 after the clock fix, check those are not setup frames.").font = NOTE_FONT
    autosize(wp, [6, 40, 10, 11, 12, 12, 12, 13, 13, 12, 12, 12, 12, 12, 12, 13, 13, 13, 11, 11, 11, 11, 11, 11])
    wp.freeze_panes = "C2"

    # ---------------- Comparison ----------------
    wc = wb.create_sheet("Comparison")
    ch = ["Pin", "Bridge", "Period", "Predicted new items (mid)", "Predicted low", "Predicted high",
          "ACTUAL new items", "ACTUAL daylight frames processed", "ACTUAL daylight hours covered",
          "Predicted items per daylight h", "Actual items per daylight h", "Actual / Predicted", "Inside low to high?", "Notes"]
    for c, h in enumerate(ch, start=1):
        wc.cell(row=1, column=c, value=h)
    style_header(wc, 1, len(ch))
    r = 2
    for name, s, e in PERIODS:
        ndays = (e - s).days + 1
        for site in SITES:
            pin = site["id"]
            p = pred.get((pin, name))
            b = per.get((pin, name))
            has_cam = any(c["id"] == pin for c in cams)
            wc.cell(row=r, column=1, value=pin)
            wc.cell(row=r, column=2, value=site["bridge"])
            wc.cell(row=r, column=3, value=name)
            wc.cell(row=r, column=4, value=p[0] if p else "")
            wc.cell(row=r, column=5, value=p[1] if p else "")
            wc.cell(row=r, column=6, value=p[2] if p else "")
            wc.cell(row=r, column=7, value=b["new"] if b else "")
            wc.cell(row=r, column=8, value=b["frames"] if b else "")
            wc.cell(row=r, column=9, value="=IF(H%d=\"\",\"\",ROUND(H%d*%d/3600,2))" % (r, r, interval))
            wc.cell(row=r, column=10, value="=IF(D%d=\"\",\"\",ROUND(D%d/%.2f,2))" % (r, r, dl_sum[name]))
            wc.cell(row=r, column=11, value="=IF(OR(G%d=\"\",H%d=\"\",H%d=0),\"\",ROUND(G%d/(H%d*%d/3600),2))" % (r, r, r, r, r, interval))
            wc.cell(row=r, column=12, value="=IF(OR(K%d=\"\",J%d=\"\",J%d=0),\"\",ROUND(K%d/J%d,2))" % (r, r, r, r, r))
            wc.cell(row=r, column=13, value="=IF(OR(K%d=\"\",E%d=\"\"),\"\",IF(AND(K%d>=E%d/%.2f,K%d<=F%d/%.2f),\"yes\",\"no\"))" % (r, r, r, r, dl_sum[name], r, r, dl_sum[name]))
            if b:
                note = "%d of %d days have frames" % (len(b["days"]), ndays)
                if b["frames"] * interval / 3600.0 < 0.5 * dl_sum[name]:
                    note += ", partial coverage"
            elif has_cam:
                note = "card has no daylight frames in this period"
            else:
                note = "no card processed for this camera yet"
            wc.cell(row=r, column=14, value=note)
            style_body(wc, r, len(ch))
            if b:
                for c in (7, 8):
                    wc.cell(row=r, column=c).fill = GREEN
            r += 1
    wc.cell(row=r + 1, column=1, value="Green cells were filled by build_field_workbook.py from detections.csv. Predicted columns come from %s." % (pred_src or "nothing, the predicted workbook was not found")).font = NOTE_FONT
    wc.cell(row=r + 2, column=1, value="Compare on items per daylight hour, because real coverage per site is not a full 14 days. Inside low to high compares the per hour rate to the scenario band. Period daylight: P1 %.1f h, P2 %.1f h, P3 %.1f h." % (dl_sum["P1"], dl_sum["P2"], dl_sum["P3"])).font = NOTE_FONT
    autosize(wc, [6, 40, 8, 16, 12, 12, 16, 20, 18, 16, 16, 14, 14, 44])
    wc.freeze_panes = "D2"

    # ---------------- Heatmap_Long ----------------
    wh = wb.create_sheet("Heatmap_Long")
    hh = ["site_id", "bridge", "lat", "lon", "period", "period_start", "period_end", "days_with_frames",
          "daylight_hours_covered", "observed_new_items", "items_per_daylight_h", "weight_0_to_1"]
    long_rows = []
    rates = []
    for name, s, e in PERIODS:
        for site in SITES:
            b = per.get((site["id"], name))
            hours_exact = b["frames"] * interval / 3600.0 if b else None
            hours = r2(hours_exact) if b else None
            rate = r2(b["new"] / hours_exact) if (b and hours_exact) else None
            if rate is not None:
                rates.append(rate)
            long_rows.append([site["id"], site["bridge"], site["lat"], site["lon"], name, s.isoformat(), e.isoformat(),
                              len(b["days"]) if b else "", hours if hours is not None else "", b["new"] if b else "", rate if rate is not None else "", None])
    mx = max(rates) if rates else 0
    for row in long_rows:
        row[11] = float(Decimal(str(row[10] / mx)).quantize(Decimal("0.001"), rounding=ROUND_HALF_UP)) if (row[10] != "" and mx) else ""
    n = write_table(wh, hh, long_rows, widths=[8, 42, 10, 11, 8, 12, 12, 10, 14, 14, 14, 12])
    wh.cell(row=n + 2, column=1, value="Long format for the heat map tool: lat, lon, weight. weight_0_to_1 is the per hour rate scaled to the busiest site period, so cameras with short coverage are not penalised. Blank rows have no data yet.").font = NOTE_FONT
    with open(os.path.join(out_dir, "heatmap_observed_long.csv"), "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(hh)
        w.writerows(long_rows)

    # ---------------- Daily ----------------
    wd = wb.create_sheet("Daily")
    dh = ["date", "period", "site_id", "bridge", "lat", "lon", "daylight_frames", "daylight_hours", "images_with_trash",
          "total_boxes", "stationary", "drift", "new_items", "new_per_daylight_h", "predicted_new_items", "flow_cfs"]
    drows = []
    for cam in cams:
        site = SITE_BY_ID.get(cam["id"], {})
        for d, b in sorted(bucket(cam, lambda r: r["date"]).items()):
            hours_exact = b["frames"] * interval / 3600.0
            hours = r2(hours_exact)
            pdaily = pred_daily.get((cam["id"], d))
            drows.append([d.isoformat(), period_of(d), cam["id"], site.get("bridge", ""), site.get("lat", ""), site.get("lon", ""),
                          b["frames"], hours, b["with_trash"], b["boxes"], b["stationary"], b["drift"], b["new"],
                          r2(b["new"] / hours_exact) if hours_exact else "",
                          pdaily[0] if pdaily else "", pdaily[1] if (pdaily and pdaily[1] is not None) else ""])
    n = write_table(wd, dh, drows, widths=[11, 8, 7, 36, 10, 11, 10, 10, 10, 10, 10, 10, 10, 12, 12, 9])
    wd.cell(row=n + 2, column=1, value="One row per camera per day. predicted_new_items and flow_cfs come from heatmap_daily_long.csv next to the predicted workbook, blank if that file was not found.").font = NOTE_FONT
    with open(os.path.join(out_dir, "daily_observed_long.csv"), "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(dh)
        w.writerows(drows)

    # ---------------- Validation ----------------
    wv = wb.create_sheet("Validation")
    vh = ["Pin", "Bridge", "Positives reviewed", "Confirmed trash", "False alarms", "Precision", "Precision CI low", "Precision CI high",
          "Negatives reviewed", "Trash the model missed", "Miss rate", "Miss rate CI low", "Miss rate CI high",
          "Recall", "Recall CI low", "Recall CI high"]
    vrows = []
    for cam in cams:
        v = cam["validation"]
        if not v:
            continue
        site = SITE_BY_ID.get(cam["id"], {})
        vrows.append([cam["id"], site.get("bridge", ""), v["positive_sample"], v["positive_confirmed"], v["positive_sample"] - v["positive_confirmed"],
                      round(v["precision"][0], 3), round(v["precision"][1], 3), round(v["precision"][2], 3),
                      v["negative_sample"], v["negative_missed"],
                      round(v["miss_rate"][0], 3), round(v["miss_rate"][1], 3), round(v["miss_rate"][2], 3),
                      round(v["recall"][0], 3), round(v["recall"][1], 3), round(v["recall"][2], 3)])
    n = write_table(wv, vh, vrows, widths=[6, 40, 10, 10, 10, 10, 10, 10, 10, 10, 10, 10, 10, 10, 10, 10])
    if not vrows:
        wv.cell(row=2, column=1, value="No camera has a scored review sheet yet. Fill in review/review_sheet.csv, run python3 score_review.py <camera>, then rebuild this workbook.").font = NOTE_FONT
    else:
        wv.cell(row=n + 2, column=1, value="From review/validation_result.txt written by score_review.py. Wilson 95 percent intervals. Fifty per stratum is about plus or minus 14 points.").font = NOTE_FONT

    # ---------------- Detections ----------------
    wdet = wb.create_sheet("Detections")
    deth = ["camera_id", "filename", "timestamp", "date", "period", "box_index", "x1", "y1", "x2", "y2", "confidence", "repeat_type", "is_new"]
    detrows = []
    for cam in cams:
        for b in cam["boxes"]:
            detrows.append([b["camera_id"], b["filename"], b["ts"].strftime("%Y-%m-%d %H:%M:%S"), b["ts"].date().isoformat(), b["period"],
                            b["box_index"], b["x1"], b["y1"], b["x2"], b["y2"], b["confidence"], b["repeat_type"], 1 if b["repeat_type"] == "new" else 0])
    for c, h in enumerate(deth, start=1):
        wdet.cell(row=1, column=c, value=h)
    style_header(wdet, 1, len(deth))
    for r, row in enumerate(detrows, start=2):
        for c, v in enumerate(row, start=1):
            wdet.cell(row=r, column=c, value=v)
    autosize(wdet, [9, 18, 19, 11, 8, 9, 7, 7, 7, 7, 10, 12, 7])
    wdet.freeze_panes = "A2"
    wdet.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(deth)), max(len(detrows) + 1, 2))

    return wb


# ---------------------------------------------------------------------------
# MAIN
# ---------------------------------------------------------------------------

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--field-dir", default="field_data", help="folder that holds the camera_XX folders")
    ap.add_argument("--out", default="South_Platte_Field_Results.xlsx", help="workbook to write")
    ap.add_argument("--predicted", default=PREDICTED_XLSX, help="the predicted dataset workbook, for the Comparison sheet")
    ap.add_argument("--interval", type=int, default=FRAME_INTERVAL_S, help="seconds between frames, for hours covered")
    ap.add_argument("--study-year", type=int, default=STUDY_YEAR, help="real year of the deployment, camera clocks said 2025")
    args = ap.parse_args()

    if not os.path.isdir(args.field_dir):
        print("No folder called %s here. Run process_camera.py on a card first, or pass --field-dir." % args.field_dir)
        return
    folders = sorted(d for d in os.listdir(args.field_dir) if re.match(r"camera_\d+$", d))
    print("Found %d camera folders in %s" % (len(folders), args.field_dir))
    cams = []
    for d in folders:
        cam = read_camera(os.path.join(args.field_dir, d), args.study_year, args.interval)
        if cam:
            cams.append(cam)
            tot = bucket(cam, lambda r: "all").get("all", {})
            print("  camera %02d: %d frames, %d night skipped, %d new items%s" % (
                cam["id"], len(cam["frames"]), cam["night"], tot.get("new", 0),
                ", clock year fixed" if cam["year_fixed"] else ""))
    if not cams:
        print("No camera folder had a detections.csv. Nothing to build.")
        return
    cams.sort(key=lambda c: c["id"])

    pred, pred_src = read_predicted(args.predicted)
    pred_daily = read_predicted_daily(args.predicted)
    print("Predicted numbers:", pred_src or "not found, Comparison will have blank predicted columns")

    out_dir = os.path.dirname(os.path.abspath(args.out))
    wb = build_workbook(cams, pred, pred_src, pred_daily, args.interval, out_dir)
    wb.save(args.out)
    print()
    print("Saved", args.out)
    print("Also wrote heatmap_observed_long.csv and daily_observed_long.csv next to it for the heat map tool.")


if __name__ == "__main__":
    main()
